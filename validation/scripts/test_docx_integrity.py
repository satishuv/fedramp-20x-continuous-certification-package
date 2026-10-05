"""Word-document integrity and sample residue: regressions for review items 1 and 5.

  AUD-F35  The authoring .docx is byte-reproducible and therefore inside every
           integrity check: regenerate-and-diff (CI, AWS buildspec, Actions
           loop), the reproducibility gate (CI and sdr.py), and the release
           manifest (so the bundle ships the checked bytes). It used to be
           excluded from all of them because python-docx stamps the save time
           into every zip entry.
  AUD-F36  Running the Class C sample as documented left the fictional offering
           in the committed Class C .docx (409 occurrences) and the validation
           reports in sample state: the restore rebuilt inactive classes' JSON
           and text but never their Word document, and CI could not see it
           because of AUD-F35. `sdr.py build` now regenerates every class
           (build_inactive_classes.py) and both sample builders restore the
           reports they overwrite.

Fast checks run always; the end-to-end sample residue check (builds and runs
package-preflight, about a minute) runs unless SDR_SKIP_SLOW_TESTS is set.

Run: python validation/scripts/test_docx_integrity.py
"""
import hashlib
import io
import os
import subprocess
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, HERE)
sys.path.insert(0, BASE)

import build_docx as bd  # noqa: E402
import build_release_manifest as brm  # noqa: E402
import build_inactive_classes as bic  # noqa: E402
import sdr  # noqa: E402

VALIDATE_WF = os.path.join(BASE, ".github", "workflows", "validate.yml")
LOOP_WF = os.path.join(BASE, ".github", "workflows", "living-sdr-loop.yml")
BUILDSPEC = os.path.join(BASE, "automation", "pipeline", "buildspec-validate.yml")
DOCX_DIR = os.path.join(BASE, "sdr", "human-readable")


def _read(path):
    with open(path, encoding="utf-8") as f:
        return f.read()


def _sha(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        h.update(f.read())
    return h.hexdigest()


def _save_docx_with_wall_clock(tmpdir, name):
    """A python-docx document saved the normal way: zip entries carry the
    current time, so two saves seconds apart differ."""
    from docx import Document
    doc = Document()
    doc.add_paragraph("Deterministic container probe")
    path = os.path.join(tmpdir, name)
    doc.save(path)
    return path


# ---- AUD-F35: byte-reproducible container ------------------------------------

def test_f35_normalize_makes_identical_content_identical_bytes():
    import tempfile
    import time
    tmp = tempfile.mkdtemp(prefix="docx-f35-")
    a = _save_docx_with_wall_clock(tmp, "a.docx")
    time.sleep(1.1)  # zip date_time has 2-second granularity; cross a boundary
    time.sleep(1.1)
    b = _save_docx_with_wall_clock(tmp, "b.docx")
    # Before normalization the containers differ (the wall clock leaked in).
    stamps_a = {i.date_time for i in zipfile.ZipFile(a).infolist()}
    stamps_b = {i.date_time for i in zipfile.ZipFile(b).infolist()}
    assert stamps_a != stamps_b, "probe did not cross a zip timestamp boundary"
    bd.normalize_docx(a)
    bd.normalize_docx(b)
    assert _sha(a) == _sha(b), "normalized containers must be byte-identical"
    infos = zipfile.ZipFile(a).infolist()
    assert {i.date_time for i in infos} == {bd.ZIP_FIXED_DATETIME}
    assert {i.create_system for i in infos} == {3}
    assert {i.external_attr for i in infos} == {0o644 << 16}
    # Still a valid Word document: python-docx can open the normalized file.
    from docx import Document
    assert Document(a).paragraphs[0].text == "Deterministic container probe"


def test_f35_document_timestamp_is_dataset_derived_not_wall_clock():
    import datetime as dt
    assert bd.document_timestamp("2026.09.13.02") == dt.datetime(2026, 9, 13)
    assert bd.document_timestamp("garbage") == bd.document_timestamp(None)
    assert bd.document_timestamp("2026.13.40.01") == bd.document_timestamp(None)  # invalid date -> fallback
    src = _read(os.path.join(HERE, "build_docx.py"))
    assert "normalize_docx(out)" in src, "build_docx must normalize what it saves"
    assert "cp.created = stamp" in src and "cp.modified = stamp" in src
    assert "datetime.now" not in src and "date.today" not in src


def test_f35_committed_docx_are_normalized():
    for cls in ("a", "b", "c"):
        path = os.path.join(DOCX_DIR, f"sdr-class-{cls}-authoring.docx")
        infos = zipfile.ZipFile(path).infolist()
        assert {i.date_time for i in infos} == {bd.ZIP_FIXED_DATETIME}, path


def test_f35_docx_is_inside_every_integrity_check():
    # Release manifest fingerprints the active-class Word document.
    assert any(g.endswith("-authoring.docx") for g in brm.ARTIFACT_GLOBS), brm.ARTIFACT_GLOBS
    # sdr.py reproducibility includes it.
    src = _read(os.path.join(BASE, "sdr.py"))
    repro = src[src.find("def cmd_reproducibility"):src.find("def cmd_reproducibility") + 1800]
    assert '".docx"' in repro and 'not fn.endswith(".docx")' not in repro
    # CI: no docx exclusion in the diff, docx in both reproducibility finds.
    wf = _read(VALIDATE_WF)
    assert "':(exclude)*.docx'" not in wf
    assert wf.count("-name '*.docx'") == 2, "both reproducibility finds must include .docx"
    assert "! -name '*.docx'" not in wf
    # AWS buildspec and Actions loop: no docx exclusion.
    assert "':(exclude)*.docx'" not in _read(BUILDSPEC)
    assert "':(exclude)*.docx'" not in _read(LOOP_WF)


def test_f35_manifest_fingerprints_the_committed_docx_bytes():
    import json
    manifest = json.load(open(os.path.join(BASE, "artifacts", "release-manifest.json"), encoding="utf-8"))
    cls = (manifest.get("certification_class") or "b").lower()
    rel = f"sdr/human-readable/sdr-class-{cls}-authoring.docx"
    assert rel in manifest["artifacts"], sorted(manifest["artifacts"])
    assert manifest["artifacts"][rel] == "sha256:" + _sha(os.path.join(BASE, rel.replace("/", os.sep)))


# ---- AUD-F36: one build definition, no sample residue ------------------------

def test_f36_build_regenerates_inactive_classes_before_the_active_one():
    steps = [s for s, _d in sdr.BUILD_STEPS]
    assert "build_inactive_classes.py" in steps
    assert steps.index("build_inactive_classes.py") < steps.index("build_sdr.py")
    assert steps.index("build_inactive_classes.py") > steps.index("build_profiles.py")
    assert bic.PER_CLASS_STEPS == ("build_sdr.py", "build_docx.py")
    assert bic.inactive_classes("b") == ["a", "c"]
    assert bic.inactive_classes("c") == ["a", "b"]


def test_f36_sample_builders_restore_the_reports_and_rebuild_everything():
    c = _read(os.path.join(BASE, "examples", "sample-offering-class-c", "build_class_c_sample.py"))
    assert "validation-report.json" in c and "ksi-test-results.json" in c
    import re
    calls = re.findall(r"^\s+_restore_tree\(baks, history_existed, tmp\)", c, flags=re.M)
    assert len(calls) == 2, "both finally blocks must share the restore"
    assert 'SDR_BUILD_CLASS=c' not in c, "per-class rebuild now lives in sdr.py build"
    b = _read(os.path.join(BASE, "examples", "sample-offering", "build_sample.py"))
    assert "validation-report.json" in b and "ksi-test-results.json" in b
    assert "for real, b in report_baks.items():" in b


def test_f36_running_the_class_c_sample_leaves_the_tree_byte_identical():
    """End to end: the exact reproduction from the delivery review. Every
    committed deliverable, the three Word documents and the validation reports
    included, must be byte-identical after the sample run, and the Class C
    document must carry none of the fictional offering."""
    if os.environ.get("SDR_SKIP_SLOW_TESTS"):
        print("  (skipped end-to-end sample run: SDR_SKIP_SLOW_TESTS set)")
        return
    watched = []
    for root in ("sdr", "package", os.path.join("validation", "reports")):
        for dirpath, _dirs, files in os.walk(os.path.join(BASE, root)):
            for fn in files:
                watched.append(os.path.join(dirpath, fn))
    watched.append(os.path.join(BASE, "artifacts", "release-manifest.json"))
    before = {p: _sha(p) for p in watched}
    r = subprocess.run([sys.executable, os.path.join(BASE, "examples", "sample-offering-class-c",
                                                     "build_class_c_sample.py")],
                       cwd=BASE, capture_output=True, text=True)
    assert r.returncode == 0, r.stdout[-2000:] + r.stderr[-2000:]
    after = {p: _sha(p) for p in watched}
    changed = sorted(os.path.relpath(p, BASE) for p in watched if before[p] != after.get(p))
    assert not changed, "sample run changed committed files: " + ", ".join(changed[:10])
    xml = zipfile.ZipFile(os.path.join(DOCX_DIR, "sdr-class-c-authoring.docx")).read("word/document.xml")
    assert b"Beacon Federal Cloud" not in xml and b"[SAMPLE" not in xml


def _run_all():
    tests = [v for k, v in sorted(globals().items())
             if k.startswith("test_") and callable(v)]
    for t in tests:
        t()
        print(f"PASS: {t.__name__}")
    print(f"\n{len(tests)}/{len(tests)} passed")
    return 0


if __name__ == "__main__":
    sys.exit(_run_all())
