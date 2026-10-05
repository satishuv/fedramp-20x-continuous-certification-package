# -*- coding: utf-8 -*-
# Generate the authoring-format SDR as a Word document for the class selected
# in the offering profile. This is the human working view a consultant fills in
# with the customer; the JSON produced by build_sdr.py remains the record.
#
# Every rule and KSI entry carries, before the fill-in fields:
#   What it looks for      the official statement resolved for this class
#   Official notes         the dataset's own notes where present
#   How to comply          SAS advisory guidance, clearly labeled
#   Evidence required      rule-specific artifacts, else FedRAMP defaults
#
# Pipeline position: after build_notes.py and build_profiles.py.
# Output: sdr/human-readable/sdr-class-<x>-authoring.docx
#
# Byte-reproducible (AUD-F35). python-docx stamps every zip entry with the save
# time, so two builds of unchanged inputs used to differ and the .docx had to be
# excluded from the regenerate-and-diff, the reproducibility gate and the
# release manifest: the one document an assessor reads was outside every
# integrity check. normalize_docx() rewrites the container with a fixed entry
# timestamp, fixed attributes and a fixed compression level, and the core
# properties carry a timestamp derived from the pinned dataset version instead
# of the wall clock, so the Word document is now fingerprinted and diffed like
# every other deliverable.

import datetime as _dt
import json
import os
import re
import sys
import zipfile

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
except ImportError:
    sys.exit(
        "build_docx.py requires the python-docx package. Install it with:\n"
        "  pip install python-docx\n"
        "(the package name is python-docx; the unrelated PyPI package named "
        "'docx' will not work)"
    )

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# Zip entry timestamp for every member: the DOS epoch, the conventional value
# for reproducible archives. Nothing about the document depends on it.
ZIP_FIXED_DATETIME = (1980, 1, 1, 0, 0, 0)
ZIP_COMPRESSLEVEL = 6


def normalize_docx(path):
    """Rewrite the .docx container in place so identical content yields
    identical bytes: member order preserved, fixed per-entry timestamp,
    attributes and creator system, fixed deflate level. Returns the path."""
    with zipfile.ZipFile(path) as zin:
        members = [(info.filename, zin.read(info.filename)) for info in zin.infolist()]
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED,
                         compresslevel=ZIP_COMPRESSLEVEL) as zout:
        for name, data in members:
            info = zipfile.ZipInfo(name, date_time=ZIP_FIXED_DATETIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3          # fixed, not the host OS
            info.external_attr = 0o644 << 16
            zout.writestr(info, data, compresslevel=ZIP_COMPRESSLEVEL)
    os.replace(tmp, path)
    return path


def document_timestamp(dataset_version):
    """A deterministic core-properties timestamp: the pinned dataset's release
    date (YYYY.MM.DD prefix of the version string), midnight UTC; a fixed
    fallback if the version does not carry a date."""
    m = re.match(r"^(\d{4})\.(\d{2})\.(\d{2})", str(dataset_version or ""))
    if m:
        try:
            return _dt.datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            pass
    return _dt.datetime(2026, 1, 1)

FILL = "[[ FILL: replace with the provider's real implementation facts ]]"


def load(*parts):
    with open(os.path.join(BASE, *parts), encoding="utf-8") as f:
        return json.load(f)


def label_para(doc, label, text, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(label + " ")
    r.bold = True
    r2 = p.add_run(text)
    r2.italic = italic
    return p


def main():
    profile = load("profiles", "common", "offering-profile.json")
    cls = (os.environ.get("SDR_BUILD_CLASS") or profile["certification_class"]).lower()
    if cls not in ("a", "b", "c"):
        print("Class D is FedRAMP pending; no authoring document generated.")
        return 1
    class_profile = load("profiles", f"class-{cls}", "profile.json")
    rules = class_profile["rules"]
    ksis = load("profiles", "common", "ksi-profile.json")["indicators"]
    if cls == "a":
        tier = class_profile["meta"]["class_a_ksis"]
        ksis = [k for k in ksis if k["ksi_id"] in tier]
    rule_notes = load("traceability", "rule-notes.json")["rules"]
    ksi_notes = load("traceability", "ksi-notes.json")["indicators"]
    records = load("sdr", "records", "records-store.json")
    fam_names = load("traceability", "family-names.json")

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    doc.add_heading("Security Decision Record, authoring document", level=0)
    doc.add_paragraph(
        f"Offering: {profile['offering_name']} ({profile['offering_abbreviation']})")
    doc.add_paragraph(
        f"Certification: FedRAMP 20x, Class {cls.upper()}, "
        f"{profile['certification_path']} path")
    doc.add_paragraph(
        f"AWS partition: {profile['aws_partition']} | regions: "
        f"{profile['primary_region']}, {profile['dr_region']}")
    doc.add_paragraph(
        f"Generated from dataset {class_profile['meta']['dataset_version']}")

    doc.add_heading("How to use this document", level=1)
    doc.add_paragraph(
        "This is the human authoring view of the Security Decision Record. "
        "Work through every entry: read the note that explains what the "
        "requirement looks for and what evidence FedRAMP expects, then replace "
        "the FILL marker with the provider's real implementation, validation, "
        "and owner. The JSON generated by the pipeline is the record of "
        "submission; this document is where the content is authored and "
        "reviewed. Guidance lines are SAS advisory content, not FedRAMP text.")
    doc.add_paragraph(
        f"This class requires {'0' if cls == 'a' else ('at least 1' if cls == 'b' else 'at least 2')} "
        "automated verification and validation method(s) per Key Security "
        "Indicator (rule FRC-CSX-VVK)."
        + ("" if cls == "a" else
           " Class C additionally requires 6 months of historical metrics "
           "before application (rule FRC-CSX-MOT)." if cls == "c" else ""))

    doc.add_heading(f"Part 1: FedRAMP rules ({len(rules)} entries)", level=1)
    for r in rules:
        rid = r["rule_id"]
        note = rule_notes.get(rid, {})
        rec = records["frr"].get(rid, {})
        doc.add_heading(f"{rid} {r['name'] or ''}", level=2)
        fam = r["family"]
        fam_full = fam_names["frr"].get(fam, fam)
        doc.add_paragraph(f"Rule family: {fam} ({fam_full})")
        force = r.get("force") or "stated in rule text"
        line = f"Force at Class {cls.upper()}: {force}"
        if r.get("class_a_obligation"):
            line += f" | Class A obligation: {r['class_a_obligation']}"
        if r.get("timeframe_num"):
            line += f" | timeframe: {r['timeframe_num']} {r.get('timeframe_type') or ''}"
        elif r.get("timeframe_num_min") is not None or r.get("timeframe_num_max") is not None:
            lo, hi = r.get("timeframe_num_min"), r.get("timeframe_num_max")
            rng = (f"{lo}-{hi}" if lo is not None and hi is not None
                   else str(lo if lo is not None else hi))
            line += f" | timeframe: {rng} {r.get('timeframe_type') or ''}"
        doc.add_paragraph(line)
        label_para(doc, "What it looks for:", r["statement"])
        for n in note.get("official_notes", []):
            label_para(doc, "Official note:", n, italic=True)
        label_para(doc, "How to comply (SAS guidance):",
                   note.get("how_to_comply_guidance", ""))
        ev = note.get("evidence_required", [])
        label_para(doc, "Evidence required "
                   f"({note.get('evidence_source', '')}):", "")
        for e in ev:
            doc.add_paragraph(str(e), style="List Bullet")
        imp = rec.get("implementation", [FILL])
        val = rec.get("validation", [FILL])
        ext = rec.get("extension", {})
        label_para(doc, "Implementation:",
                   imp[0] if imp and "TBD" not in imp[0] else FILL)
        label_para(doc, "Validation:",
                   val[0] if val and "TBD" not in val[0] else FILL)
        label_para(doc, "Owner:",
                   ext.get("owner", FILL) if "TBD" not in ext.get("owner", "TBD") else FILL)
        label_para(doc, "Status:", rec.get("implementation_status", "Not Implemented"))

    doc.add_heading(f"Part 2: Key Security Indicators ({len(ksis)} entries)", level=1)
    for k in ksis:
        kid = k["ksi_id"]
        note = ksi_notes.get(kid, {})
        rec = records["ksi"].get(kid, {})
        doc.add_heading(f"{kid} {k['name'] or ''}", level=2)
        doc.add_paragraph(f"Family: {k['family']} ({k['family_name']})")
        label_para(doc, "What it looks for:", note.get("what_it_looks_for", ""))
        label_para(doc, "How to comply (SAS guidance):",
                   note.get("how_to_comply_guidance", ""))
        controls = note.get("nist_controls", [])
        if controls:
            label_para(doc, "Related NIST SP 800-53 Rev5 controls:",
                       ", ".join(controls))
        label_para(doc, "Evidence required (FedRAMP defaults for KSIs):", "")
        for e in note.get("evidence_required", []):
            doc.add_paragraph(str(e), style="List Bullet")
        label_para(doc, "Automated methods required at this class:",
                   str(k["minimum_automated_methods"][f"class_{cls}"]))
        label_para(doc, "Historical metrics required:",
                   k["historical_metrics"][f"class_{cls}"])
        imp = rec.get("implementation", [FILL])
        val = rec.get("validation", [FILL])
        ext = rec.get("extension", {})
        label_para(doc, "Implementation:",
                   imp[0] if imp and "TBD" not in imp[0] and "FedRAMP pending" not in imp[0] else FILL)
        label_para(doc, "Validation:",
                   val[0] if val and "TBD" not in val[0] else FILL)
        tests = rec.get("tests", [])
        # A test may be a plain string or a structured record ({method,
        # automated, cadence}); coerce so a structured entry does not crash.
        def _t(t):
            if isinstance(t, dict):
                m = t.get("method") or t.get("name") or t.get("description") or "test"
                extra = []
                if t.get("automated") is True:
                    extra.append("automated")
                elif t.get("automated") is False:
                    extra.append("manual")
                if t.get("cadence"):
                    extra.append(str(t["cadence"]))
                return m + (f" ({', '.join(extra)})" if extra else "")
            return str(t)
        label_para(doc, "Tests:",
                   "; ".join(_t(t) for t in tests) if tests else FILL)
        label_para(doc, "Owner:",
                   ext.get("owner", FILL) if "TBD" not in ext.get("owner", "TBD") else FILL)
        label_para(doc, "Status:", rec.get("implementation_status", "Not Implemented"))

    doc.add_heading("Completeness self-check before assessment", level=1)
    checks = [
        f"Every one of the {len(rules)} rule entries has real implementation text, no FILL markers remain.",
        f"Every one of the {len(ksis)} indicator entries has real implementation text and a named owner.",
        "Every indicator meets the automated method minimum for this class (FRC-CSX-VVK).",
        "The JSON SDR regenerated by the pipeline validates with zero schema errors.",
        "Human-readable and JSON forms were generated from the same record store in the same run.",
    ]
    if cls == "c":
        checks.append("At least 6 months of daily metrics history exists for every indicator (FRC-CSX-MOT).")
    for c in checks:
        doc.add_paragraph(c, style="List Number")

    cp = doc.core_properties
    cp.author = "SDR framework generator"
    cp.title = f"Security Decision Record authoring document, Class {cls.upper()}"
    # AUD-F35: no wall-clock value anywhere in the document. The timestamp is
    # the pinned dataset's release date, so an unchanged input set produces an
    # unchanged document, byte for byte.
    stamp = document_timestamp(class_profile["meta"].get("dataset_version"))
    cp.created = stamp
    cp.modified = stamp
    cp.last_printed = stamp
    cp.last_modified_by = "SDR framework generator"
    cp.revision = 1

    out = os.path.join(BASE, "sdr", "human-readable", f"sdr-class-{cls}-authoring.docx")
    doc.save(out)
    normalize_docx(out)
    print("saved:", out, "(byte-reproducible container)")
    print("rules:", len(rules), "| ksis:", len(ksis))
    return 0


if __name__ == "__main__":
    sys.exit(main())
